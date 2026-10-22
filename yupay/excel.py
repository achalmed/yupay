"""La liquidación de una rendición en Excel (2026-10-18), con las MISMAS fórmulas del libro con que el
despacho rindió la SR-01: hoja «Rendición» (un día por columna: A con comprobante, B con declaración
jurada, C liquidación diaria, D controles «CONFORME/REVISAR»), «Resumen» y «Parámetros». Los importes de
cada día van como valores (celdas azules, editables) y todo lo demás son fórmulas: quien la abra puede
probar otro monto y el libro recalcula, como el original. La fuente de verdad sigue siendo Qillqa.

Requiere `openpyxl` (requirements.txt). Pura salvo la escritura del archivo: `libro(vista)`
devuelve el Workbook; la prueba recalcula sus fórmulas en `tests/test_rendicion.py`.
"""
from __future__ import annotations

from io import BytesIO

from yupay import rendicion as r

FORMATO_SOLES = '"S/ "#,##0.00;"(S/ "#,##0.00\\);\\-'
VINO, PAPEL, BORDE, TINTA = "4E0A18", "F5F1EC", "DCD2C6", "241F1C"
AZUL = "1F4E9A"   # entradas (convención del libro de referencia: azul = lo escribe la persona)


def libro(vista: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as col

    datos = vista.get("datos") or {}
    liq = vista["liquidacion"]
    dias = [d["fecha"] for d in liq["dias"]]
    n = len(dias)
    wb = Workbook()
    ws = wb.active
    ws.title = "Rendición"
    par = wb.create_sheet("Parámetros")
    res = wb.create_sheet("Resumen")
    wb.move_sheet("Resumen", offset=-1)

    fino = Side(style="thin", color=BORDE)
    caja = Border(left=fino, right=fino, top=fino, bottom=fino)
    cab = Font(bold=True, color="FFFFFF")
    relleno_cab = PatternFill("solid", fgColor=VINO)
    relleno_sec = PatternFill("solid", fgColor=PAPEL)
    entrada = Font(color=AZUL)
    negrita = Font(bold=True, color=TINTA)

    # --- Parámetros (B3 tope, B4 tope DJ, B5 días, B6 asignado, B7 % DJ, B8 recibido) ------------
    p = {k: (v["valor"] if isinstance(v, dict) else v) for k, v in (vista.get("parametros") or {}).items()}
    par["A1"] = "PARÁMETROS DE CÁLCULO"
    par["A1"].font = Font(bold=True, size=13, color=VINO)
    filas_par = [
        ("Tope de viático diario", float(liq["tope_diario"]), (vista.get("parametros") or {}).get("tope_diario", {}).get("fuente", "")),
        ("Tope diario sustentado con declaración jurada", float(liq["tope_dj"]), (vista.get("parametros") or {}).get("tope_dj", {}).get("fuente", "")),
        ("Días de la comisión", n, ", ".join(r.fecha_corta(d) for d in dias)),
        ("Viático total asignado", "=B3*B5", "Calculado: tope diario × número de días."),
        ("% máximo de declaración jurada sobre el viático", "=B4/B3", "Calculado: tope DJ / tope diario."),
        ("Importe recibido de viáticos", float(liq["recibido"]), "Del comprobante de pago (sin dato: tope × días)."),
        ("Importe recibido de pasajes terrestres", float(p.get("recibido_pasajes") or 0), ""),
        ("Importe recibido de otros recursos", float(p.get("recibido_otros") or 0), ""),
    ]
    for i, (rotulo, valor, nota) in enumerate(filas_par, 3):
        par.cell(i, 1, rotulo)
        c = par.cell(i, 2, valor)
        c.number_format = "0.00%" if i == 7 else ("0" if i == 5 else FORMATO_SOLES)
        if not (isinstance(valor, str) and valor.startswith("=")):
            c.font = entrada
        par.cell(i, 3, nota)
    par.column_dimensions["A"].width = 46
    par.column_dimensions["B"].width = 16
    par.column_dimensions["C"].width = 90

    # --- Rendición -------------------------------------------------------------------------------
    T = col(n + 2)            # columna TOTAL
    S = col(n + 3)            # columna SUSTENTO
    ult = col(n + 1)          # último día
    ws["A1"] = "RENDICIÓN DE VIÁTICOS"
    ws["A1"].font = Font(bold=True, size=13, color="FFFFFF")
    ws["A1"].fill = relleno_cab
    ws.merge_cells(f"A1:{S}1")
    ws["A2"] = (f"Del {r.fecha_corta(datos.get('fecha_inicio') or '')} al {r.fecha_corta(datos.get('fecha_fin') or '')}  ·  "
                f"Tope de viático {r.texto_soles(r.centimos(liq['tope_diario']))} por día  ·  Declaración jurada hasta "
                f"{r.texto_soles(r.centimos(liq['tope_dj']))} por día")
    ws.merge_cells(f"A2:{S}2")
    encabezado = ["CONCEPTO"] + [f"Día {d[8:10]}/{d[5:7]}" for d in dias] + ["TOTAL", "SUSTENTO"]
    for j, t in enumerate(encabezado, 1):
        c = ws.cell(4, j, t)
        c.font, c.fill, c.border = cab, relleno_cab, caja
        c.alignment = Alignment(horizontal="center")

    def seccion(fila: int, texto: str) -> None:
        ws.cell(fila, 1, texto).font = negrita
        for j in range(1, n + 4):
            ws.cell(fila, j).fill = relleno_sec

    def fila_valores(fila: int, rotulo: str, valores: list[float], sustento: str) -> None:
        ws.cell(fila, 1, rotulo)
        for j, v in enumerate(valores, 2):
            c = ws.cell(fila, j, v if v else None)
            c.number_format, c.font = FORMATO_SOLES, entrada
        ws.cell(fila, n + 2, f"=SUM(B{fila}:{ult}{fila})").number_format = FORMATO_SOLES
        ws.cell(fila, n + 3, sustento)

    def fila_formula(fila: int, rotulo: str, formula, sustento: str, total: str | None = None, fuerte: bool = False) -> None:
        ws.cell(fila, 1, rotulo).font = negrita if fuerte else Font(color=TINTA)
        for j in range(2, n + 2):
            c = ws.cell(fila, j, formula(col(j)))
            c.number_format = FORMATO_SOLES
            if fuerte:
                c.font = negrita
        t = ws.cell(fila, n + 2, total or f"=SUM(B{fila}:{ult}{fila})")
        t.number_format, t.font = FORMATO_SOLES, negrita
        ws.cell(fila, n + 3, sustento)

    seccion(5, "A.  GASTOS CON COMPROBANTE DE PAGO (factura o boleta)")
    for i, concepto in enumerate(r.CONCEPTOS_A, 6):
        fila_valores(i, r.CONCEPTOS[concepto][1], [d["por_concepto"][concepto] for d in liq["dias"]], "Factura / boleta")
    fila_formula(11, "SUBTOTAL CON COMPROBANTE", lambda c: f"=SUM({c}6:{c}10)", "Suma de la sección A", fuerte=True)
    seccion(13, f"B.  GASTOS SUSTENTADOS CON DECLARACIÓN JURADA (tope {r.texto_soles(r.centimos(liq['tope_dj']))} por día)")
    fila_valores(14, "Movilidad local (declarada)", [d["movilidad"] for d in liq["dias"]],
                 "Menor entre el tope DJ y lo que falta para completar el viático del día.")
    fila_valores(15, "Otros gastos sin comprobante", [d["otros_dj"] for d in liq["dias"]], "Tabla OTROS del Anexo N.° 3.")
    fila_formula(16, "SUBTOTAL DECLARACIÓN JURADA", lambda c: f"=SUM({c}14:{c}15)", "Detalle de tramos del Anexo N.° 3.", fuerte=True)
    seccion(18, "C.  LIQUIDACIÓN DIARIA")
    fila_formula(19, "Viático asignado por día", lambda c: "=Parámetros!$B$3", "Tope diario autorizado")
    fila_formula(20, "Total rendido (reconocido)", lambda c: f"=MIN({c}11,Parámetros!$B$3)+{c}16", "Comprobantes (hasta el tope) + declaración jurada")
    fila_formula(21, "Exceso no reconocido", lambda c: f"=MAX(0,{c}11-Parámetros!$B$3)", "Gasto sobre el tope diario: lo asume el comisionado")
    fila_formula(22, "SALDO A DEVOLVER", lambda c: f"={c}19-{c}20", "Monto a reintegrar a la entidad", fuerte=True)
    seccion(24, "D.  CONTROLES AUTOMÁTICOS")
    controles = [
        (25, "¿Declaración jurada dentro del tope?", lambda c: f'=IF({c}16<=Parámetros!$B$4,"OK","EXCEDE")'),
        (26, "¿Total del día dentro del tope?", lambda c: f'=IF({c}20<=Parámetros!$B$3+0.001,"OK","EXCEDE")'),
        (27, "¿Existe gasto no reconocido?", lambda c: f'=IF({c}21>0,"REVISAR","OK")'),
    ]
    for fila, rotulo, formula in controles:
        ws.cell(fila, 1, rotulo)
        for j in range(2, n + 2):
            ws.cell(fila, j, formula(col(j))).alignment = Alignment(horizontal="center")
        ws.cell(fila, n + 2, f'=IF(COUNTIF(B{fila}:{ult}{fila},"OK")={n},"CONFORME","REVISAR")').font = negrita
    notas = [
        "NOTAS",
        "1.  Los topes y lo recibido se cambian en la hoja «Parámetros», no dentro de las fórmulas.",
        "2.  Celdas AZULES = importes de los comprobantes y de la declaración jurada (editables). Las demás son fórmulas.",
        "3.  La movilidad se declara al máximo permitido: el menor entre el tope DJ y lo que falta para completar el viático del día.",
        "4.  El detalle de tramos del Anexo N.° 3 sustenta, día por día, la fila 16. Lo que pase del tope lo asume el comisionado.",
        "5.  Generado por Qillqa (módulo Rendición de cuentas); la fuente de verdad son sus datos.",
    ]
    for i, t in enumerate(notas, 29):
        ws.cell(i, 1, t).font = negrita if i == 29 else Font(color=TINTA)
        ws.merge_cells(f"A{i}:{S}{i}")
    ws.column_dimensions["A"].width = 36
    for j in range(2, n + 3):
        ws.column_dimensions[col(j)].width = 13
    ws.column_dimensions[S].width = 48
    ws.freeze_panes = "B5"

    # --- Resumen ---------------------------------------------------------------------------------
    res["A1"] = f"RESUMEN DE LA COMISIÓN — {n} DÍA(S)"
    res["A1"].font = Font(bold=True, size=13, color=VINO)
    filas_res = [
        (4, "Viático total asignado (días × tope)", "=Parámetros!$B$6"),
        (5, "Gastos con comprobante (factura / boleta)", f"=Rendición!{T}11"),
        (6, "Declaración jurada (movilidad y otros)", f"=Rendición!{T}16"),
        (7, "TOTAL RENDIDO", "=B5+B6"),
        (8, "Exceso no reconocido (lo asume el comisionado)", f"=Rendición!{T}21"),
        (9, "TOTAL RECONOCIDO", "=B7-B8"),
        (10, "Importe recibido de viáticos", "=Parámetros!$B$8"),
        (11, "SALDO A DEVOLVER", "=B10-B9"),
        (12, "% de declaración jurada sobre el viático", "=B6/B4"),
    ]
    for fila, rotulo, formula in filas_res:
        res.cell(fila, 1, rotulo).font = negrita if rotulo.isupper() else Font(color=TINTA)
        c = res.cell(fila, 2, formula)
        c.number_format = "0.00%" if fila == 12 else FORMATO_SOLES
    res["C12"] = '=IF(B12<=Parámetros!$B$7+0.0001,"Dentro del límite","EXCEDE EL LÍMITE")'
    res["A14"] = "DETALLE POR DÍA"
    res["A14"].font = negrita
    for j, t in enumerate(["Día", "Con comprobante", "Declaración jurada", "Total reconocido", "Saldo a devolver", "Estado"], 1):
        c = res.cell(15, j, t)
        c.font, c.fill = cab, relleno_cab
    for i, d in enumerate(dias):
        fila, c = 16 + i, col(2 + i)
        res.cell(fila, 1, r.fecha_corta(d))
        for j, ref in enumerate((11, 16, 20, 22), 2):
            res.cell(fila, j, f"=Rendición!{c}{ref}").number_format = FORMATO_SOLES
        res.cell(fila, 6, f"=Rendición!{c}26")
    total = 16 + n
    res.cell(total, 1, "TOTAL").font = negrita
    for j in range(2, 6):
        cc = res.cell(total, j, f"=SUM({col(j)}16:{col(j)}{total - 1})")
        cc.number_format, cc.font = FORMATO_SOLES, negrita
    res.column_dimensions["A"].width = 46
    for j in range(2, 7):
        res.column_dimensions[col(j)].width = 18
    return wb


def bytes_de(vista: dict) -> bytes:
    salida = BytesIO()
    libro(vista).save(salida)
    return salida.getvalue()
